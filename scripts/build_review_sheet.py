"""
產生給受訪醫師的審閱表：data/review/病灶用藥對照_審閱表.xlsx

內容取自程式實際使用的對照表（不修改它們）：
- backend/app/data/conditions.json：病灶、關鍵字、ICD 代碼、ATC 藥理類別、常用成分
- backend/app/data/risk_flags.json：病史風險提醒
- data/drug_reference/drug_reference.sqlite：每個類別目前有有效許可證的成分
- data/raw/icd/icd10cm_2023_tw.csv：健保署 2023 年版中文 ICD-10-CM（data.gov.tw 177507），用來核對代碼

用法：python3 scripts/build_review_sheet.py
"""

from __future__ import annotations

import csv
import json
import sqlite3
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
APP_DATA = ROOT / "backend" / "app" / "data"
DRUG_DB = ROOT / "data" / "drug_reference" / "drug_reference.sqlite"
ICD_CSV = ROOT / "data" / "raw" / "icd" / "icd10cm_2023_tw.csv"
OUT = ROOT / "data" / "review" / "病灶用藥對照_審閱表.xlsx"

# WHO ATC 官方英文名稱（第 3、4 層）；中文為本專案的參考翻譯
ATC_NAMES = {
    "A01A": ("Stomatological preparations", "口腔用藥"),
    "A01AC": ("Corticosteroids for local oral treatment", "口腔局部用類固醇"),
    "A01AD": ("Other agents for local oral treatment", "其他口腔局部用藥"),
    "A02A": ("Antacids", "制酸劑"),
    "A02AD": ("Combinations and complexes of aluminium, calcium and magnesium compounds", "鋁鈣鎂複方制酸劑"),
    "A02B": ("Drugs for peptic ulcer and gastro-oesophageal reflux disease (GORD)", "消化性潰瘍及胃食道逆流用藥"),
    "A02BA": ("H2-receptor antagonists", "H2 受體阻斷劑"),
    "A02BC": ("Proton pump inhibitors", "氫離子幫浦抑制劑（PPI）"),
    "A03A": ("Drugs for functional gastrointestinal disorders", "功能性腸胃疾病用藥"),
    "A03AX": ("Other drugs for functional gastrointestinal disorders", "其他功能性腸胃疾病用藥"),
    "A03F": ("Propulsives", "腸胃蠕動促進劑"),
    "A03FA": ("Propulsives", "腸胃蠕動促進劑"),
    "A04A": ("Antiemetics and antinauseants", "止吐劑"),
    "A04AD": ("Other antiemetics", "其他止吐劑"),
    "A06A": ("Drugs for constipation", "便秘用藥"),
    "A06AA": ("Softeners, emollients", "軟便劑"),
    "A06AB": ("Contact laxatives", "刺激性瀉劑"),
    "A06AD": ("Osmotically acting laxatives", "滲透性瀉劑"),
    "A07B": ("Intestinal adsorbents", "腸道吸附劑"),
    "A07BC": ("Other intestinal adsorbents", "其他腸道吸附劑"),
    "A07D": ("Antipropulsives", "腸蠕動抑制劑"),
    "A07DA": ("Antipropulsives", "腸蠕動抑制劑"),
    "A07F": ("Antidiarrheal microorganisms", "止瀉益生菌"),
    "A07FA": ("Antidiarrheal microorganisms", "止瀉益生菌"),
    "A09A": ("Digestives, incl. enzymes", "消化劑（含酵素）"),
    "A09AA": ("Enzyme preparations", "酵素製劑"),
    "A10B": ("Blood glucose lowering drugs, excl. insulins", "降血糖藥（不含胰島素）"),
    "A10BA": ("Biguanides", "雙胍類"),
    "A10BB": ("Sulfonylureas", "磺醯脲類"),
    "A10BH": ("Dipeptidyl peptidase 4 (DPP-4) inhibitors", "DPP-4 抑制劑"),
    "A10BK": ("Sodium-glucose co-transporter 2 (SGLT2) inhibitors", "SGLT2 抑制劑"),
    "C03A": ("Low-ceiling diuretics, thiazides", "低效利尿劑（thiazide 類）"),
    "C03AA": ("Thiazides, plain", "Thiazide 類單方"),
    "C07A": ("Beta blocking agents", "乙型阻斷劑"),
    "C07AB": ("Beta blocking agents, selective", "選擇性乙型阻斷劑"),
    "C08C": ("Selective calcium channel blockers with mainly vascular effects", "血管選擇性鈣離子阻斷劑"),
    "C08CA": ("Dihydropyridine derivatives", "Dihydropyridine 類"),
    "C09A": ("ACE inhibitors, plain", "ACE 抑制劑單方"),
    "C09AA": ("ACE inhibitors, plain", "ACE 抑制劑單方"),
    "C09C": ("Angiotensin II receptor blockers (ARBs), plain", "ARB 單方"),
    "C09CA": ("Angiotensin II receptor blockers (ARBs), plain", "ARB 單方"),
    "C10A": ("Lipid modifying agents, plain", "調血脂藥單方"),
    "C10AA": ("HMG CoA reductase inhibitors", "Statin 類"),
    "C10AB": ("Fibrates", "Fibrate 類"),
    "C10AX": ("Other lipid modifying agents", "其他調血脂藥"),
    "D04A": ("Antipruritics, incl. antihistamines, anesthetics, etc.", "止癢劑（含抗組織胺、麻醉劑等）"),
    "D04AA": ("Antihistamines for topical use", "外用抗組織胺"),
    "D07A": ("Corticosteroids, plain", "外用類固醇單方"),
    "D07AA": ("Corticosteroids, weak (group I)", "弱效外用類固醇"),
    "D07AB": ("Corticosteroids, moderately potent (group II)", "中效外用類固醇"),
    "D07AC": ("Corticosteroids, potent (group III)", "強效外用類固醇"),
    "D11A": ("Other dermatological preparations", "其他皮膚用藥"),
    "D11AH": ("Agents for dermatitis, excluding corticosteroids", "非類固醇皮膚炎用藥"),
    "H02A": ("Corticosteroids for systemic use, plain", "全身性類固醇單方"),
    "H02AB": ("Glucocorticoids", "糖皮質類固醇"),
    "J01A": ("Tetracyclines", "四環黴素類"),
    "J01AA": ("Tetracyclines", "四環黴素類"),
    "J01C": ("Beta-lactam antibacterials, penicillins", "盤尼西林類"),
    "J01CA": ("Penicillins with extended spectrum", "廣效盤尼西林"),
    "J01CE": ("Beta-lactamase sensitive penicillins", "對乙內醯胺酶敏感的盤尼西林"),
    "J01CR": ("Combinations of penicillins, incl. beta-lactamase inhibitors", "盤尼西林複方（含乙內醯胺酶抑制劑）"),
    "J01D": ("Other beta-lactam antibacterials", "其他乙內醯胺類抗生素"),
    "J01DB": ("First-generation cephalosporins", "第一代頭孢菌素"),
    "J01DC": ("Second-generation cephalosporins", "第二代頭孢菌素"),
    "J01F": ("Macrolides, lincosamides and streptogramins", "巨環類、林可醯胺類"),
    "J01FA": ("Macrolides", "巨環類"),
    "J01M": ("Quinolone antibacterials", "Quinolone 類抗生素"),
    "J01MA": ("Fluoroquinolones", "Fluoroquinolone 類"),
    "J01X": ("Other antibacterials", "其他抗菌劑"),
    "J01XE": ("Nitrofuran derivatives", "Nitrofuran 類"),
    "M01A": ("Anti-inflammatory and antirheumatic products, non-steroids", "非類固醇消炎止痛藥（NSAID）"),
    "M01AB": ("Acetic acid derivatives and related substances", "醋酸衍生物"),
    "M01AE": ("Propionic acid derivatives", "丙酸衍生物"),
    "M01AH": ("Coxibs", "COX-2 抑制劑"),
    "M02A": ("Topical products for joint and muscular pain", "關節肌肉疼痛外用藥"),
    "M02AA": ("Antiinflammatory preparations, non-steroids for topical use", "外用 NSAID"),
    "M03B": ("Muscle relaxants, centrally acting agents", "中樞性肌肉鬆弛劑"),
    "M03BX": ("Other centrally acting agents", "其他中樞性肌肉鬆弛劑"),
    "M04A": ("Antigout preparations", "痛風用藥"),
    "M04AA": ("Preparations inhibiting uric acid production", "抑制尿酸生成"),
    "M04AC": ("Preparations with no effect on uric acid metabolism", "不影響尿酸代謝（如 colchicine）"),
    "N02B": ("Other analgesics and antipyretics", "其他止痛退燒藥"),
    "N02BE": ("Anilides", "Anilide 類（acetaminophen）"),
    "N02C": ("Antimigraine preparations", "偏頭痛用藥"),
    "N02CC": ("Selective serotonin (5HT1) agonists", "Triptan 類"),
    "N05B": ("Anxiolytics", "抗焦慮劑"),
    "N05BA": ("Benzodiazepine derivatives", "苯二氮平類"),
    "N05C": ("Hypnotics and sedatives", "安眠鎮靜劑"),
    "N05CD": ("Benzodiazepine derivatives", "苯二氮平類"),
    "N05CF": ("Benzodiazepine related drugs", "類苯二氮平（Z-drug）"),
    "N06A": ("Antidepressants", "抗憂鬱劑"),
    "N06AB": ("Selective serotonin reuptake inhibitors", "SSRI"),
    "N07C": ("Antivertigo preparations", "抗眩暈藥"),
    "N07CA": ("Antivertigo preparations", "抗眩暈藥"),
    "R01A": ("Decongestants and other nasal preparations for topical use", "鼻用去充血劑及其他鼻用藥"),
    "R01AC": ("Antiallergic agents, excl. corticosteroids", "鼻用抗過敏劑（不含類固醇）"),
    "R01AD": ("Corticosteroids", "鼻用類固醇"),
    "R01B": ("Nasal decongestants for systemic use", "口服鼻充血解除劑"),
    "R01BA": ("Sympathomimetics", "擬交感神經劑"),
    "R02A": ("Throat preparations", "咽喉用藥"),
    "R02AA": ("Antiseptics", "咽喉消毒劑"),
    "R03A": ("Adrenergics, inhalants", "吸入型腎上腺素類"),
    "R03AC": ("Selective beta-2-adrenoreceptor agonists", "選擇性乙二型致效劑"),
    "R03AK": ("Adrenergics in combination with corticosteroids or other drugs, excl. anticholinergics", "乙二型致效劑＋類固醇複方"),
    "R03B": ("Other drugs for obstructive airway diseases, inhalants", "其他吸入型呼吸道阻塞用藥"),
    "R03BA": ("Glucocorticoids", "吸入型類固醇"),
    "R03D": ("Other systemic drugs for obstructive airway diseases", "其他全身性呼吸道阻塞用藥"),
    "R03DC": ("Leukotriene receptor antagonists", "白三烯受體拮抗劑"),
    "R05C": ("Expectorants, excl. combinations with cough suppressants", "祛痰劑"),
    "R05CA": ("Expectorants", "祛痰劑"),
    "R05CB": ("Mucolytics", "化痰劑"),
    "R05D": ("Cough suppressants, excl. combinations with expectorants", "鎮咳劑"),
    "R05DA": ("Opium alkaloids and derivatives", "鴉片生物鹼及衍生物"),
    "R05DB": ("Other cough suppressants", "其他鎮咳劑"),
    "R05X": ("Other cold preparations", "綜合感冒製劑"),
    "R06A": ("Antihistamines for systemic use", "全身性抗組織胺"),
    "R06AB": ("Substituted alkylamines", "Alkylamine 類"),
    "R06AE": ("Piperazine derivatives", "Piperazine 類"),
    "R06AX": ("Other antihistamines for systemic use", "其他全身性抗組織胺"),
    "S01A": ("Antiinfectives", "眼用抗感染劑"),
    "S01AA": ("Antibiotics", "眼用抗生素"),
    "S01AE": ("Fluoroquinolones", "眼用 Fluoroquinolone"),
    "S01B": ("Antiinflammatory agents", "眼用消炎劑"),
    "S01BA": ("Corticosteroids, plain", "眼用類固醇"),
    "S01G": ("Decongestants and antiallergics", "眼用去充血及抗過敏劑"),
    "S01GA": ("Sympathomimetics used as decongestants", "眼用去充血劑"),
    "S01GX": ("Other antiallergics", "其他眼用抗過敏劑"),
}

# ICD-10-CM 核對結果：目前代碼的問題與建議（依健保署 2023 年版中文 ICD-10-CM）
ICD_ADVICE = {
    "咳嗽": ("R05.9", "R05 在 2023 版是類目，不能申報，需選子碼", "R05.1 急性咳嗽；R05.3 慢性咳嗽"),
    "鼻竇炎": ("J01.90", "J01.9 在 2023 版是類目，不能申報，需選子碼", "J01.91 急性復發性鼻竇炎；慢性鼻竇炎為 J32.9"),
    "過敏性鼻炎": ("J30.9", "J30.4 是 WHO 版 ICD-10 的代碼，ICD-10-CM 沒有", "J30.2 其他季節性過敏性鼻炎；J30.89 其他過敏性鼻炎"),
    "蟲咬皮膚炎": ("依部位選 S 碼＋W57.XXXA", "W57 是外因碼（事故原因），不能當主診斷；ICD-10-CM 的昆蟲叮咬依部位編碼，沒有「未明示部位」可申報碼", "例：S80.861A 右側小腿昆蟲叮咬（無毒性）初期照護；丘疹性蕁麻疹可用 L28.2 其他癢疹"),
    "肌肉骨骼疼痛": ("M79.10", "M79.1 在 2023 版是類目，不能申報，需選子碼", "M79.18 其他部位肌痛症；M25.50 關節痛；M79.7 纖維肌痛"),
    "扁桃腺炎／咽喉炎": ("J02.9（咽炎）／J03.90（扁桃腺炎）", "J02.9 正確，但只涵蓋咽炎；這個病灶同時包含扁桃腺炎", "J03.90 急性扁桃腺炎"),
    "急性腹瀉": ("A09", "可申報，但 A09 指「感染性腸胃炎及大腸炎」，若只是症狀可改用 R19.7", "R19.7 腹瀉；K52.9 非傳染性胃腸炎及結腸炎"),
    "細菌性結膜炎": ("H10.9", "可申報，但 H10.9 是「結膜炎，未明示」，沒有表達細菌性", "H10.029 未明示側性之其他黏液膿性結膜炎；H10.023 雙側"),
    "過敏性結膜炎": ("H10.45", "正確（慢性過敏性結膜炎）", "急性：H10.10 未明示側性之急性過敏性結膜炎"),
    "眩暈": ("R42", "正確（頭暈及目眩）", "H81.10 良性陣發性眩暈；H81.399 其他末梢性眩暈"),
    "焦慮": ("F41.9", "正確（非特定的焦慮症）", "F41.1 廣泛性焦慮症"),
}

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="16232C")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
EXAMPLE_FILL = PatternFill("solid", fgColor="EEF3F6")
WARN_FILL = PatternFill("solid", fgColor="FCE4E4")
OK_FILL = PatternFill("solid", fgColor="E6F2EA")
THIN = Side(style="thin", color="D5D9DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def atc_label(code: str) -> str:
    en, zh = ATC_NAMES.get(code, ("", ""))
    return f"{code} {zh}（{en}）" if en else code


def load_icd() -> dict[str, tuple[str, str]]:
    rows = {}
    with open(ICD_CSV, encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        next(reader)
        for row in reader:
            rows[row[0]] = (row[1], row[3])  # USE（1=可申報）、中文名稱
    return rows


def substance_name(code: str, name: str) -> str:
    """複方在 TFDA 資料只寫 combinations，加上代碼才看得出是哪個成分的複方。"""
    return f"{code} {name}" if "combination" in name.lower() or len(code) < 7 else name


def icd_result(name: str, use: str) -> str:
    """不能申報 → 需修正；可申報但有更貼切的代碼 → 可再確認；其餘 → 正確。"""
    if use != "1":
        return "需修正"
    note = ICD_ADVICE.get(name, ("", "正確", ""))[1]
    return "正確" if note.startswith("正確") else "可再確認"


def write_table(ws, headers: list[tuple[str, int]], rows: list[list], input_cols: set[int], start_row: int = 1):
    for col, (title, width) in enumerate(headers, 1):
        cell = ws.cell(start_row, col, title)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, values in enumerate(rows, start_row + 1):
        for col, value in enumerate(values, 1):
            cell = ws.cell(r, col, value)
            cell.font = Font(name=FONT)
            cell.alignment = WRAP
            cell.border = BORDER
            if col in input_cols:
                cell.fill = INPUT_FILL
    ws.freeze_panes = ws.cell(start_row + 1, 2)
    ws.row_dimensions[start_row].height = 34


def add_choice(ws, col: int, first: int, last: int, choices: str):
    dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(col)}{first}:{get_column_letter(col)}{last}")


def main() -> None:
    conditions = json.loads((APP_DATA / "conditions.json").read_text(encoding="utf-8"))["conditions"]
    flags = json.loads((APP_DATA / "risk_flags.json").read_text(encoding="utf-8"))["flags"]
    icd = load_icd()
    db = sqlite3.connect(DRUG_DB)
    concepts = db.execute(
        "SELECT atc_code, name_en, CAST(active_license_count AS INTEGER) FROM concepts WHERE CAST(active_license_count AS INTEGER) > 0"
    ).fetchall()
    names = dict((c, n) for c, n, _ in concepts)

    def substances(prefix: str) -> list[tuple[str, str, int]]:
        return sorted([c for c in concepts if c[0].startswith(prefix)], key=lambda c: -c[2])

    wb = Workbook()

    # ---------- 說明 ----------
    ws = wb.active
    ws.title = "說明"
    ws.column_dimensions["A"].width = 100
    lines = [
        ("新手家醫看診助手：病灶與用藥對照審閱表", Font(name=FONT, bold=True, size=16)),
        (f"產生日期：{date.today().isoformat()}。練習專案，非臨床指引。", Font(name=FONT, color="5B6770")),
        ("", None),
        ("請您幫忙確認的事", Font(name=FONT, bold=True, size=12)),
        ("系統會依病歷主訴偵測「病灶」，再依下表把病灶對應到 ATC 藥理類別，列出候選藥品給醫師參考。這份對照表由開發者依一般用藥常識整理，需要臨床醫師確認是否合理。", None),
        ("", None),
        ("怎麼填", Font(name=FONT, bold=True, size=12)),
        ("1. 只需要填黃色欄位，其他欄位是系統目前的設定，請勿修改。", None),
        ("2.「審閱結果」欄有下拉選單：同意／需修改／刪除。選「需修改」時，請在旁邊的欄位寫下要新增或移除的類別、成分。", None),
        ("3. 每個分頁的第一列資料（淺藍底）是填寫範例，不是真實審閱結果。", None),
        ("4.「缺少的病灶」分頁可以補上家醫科門診常見、但系統還沒有的病灶。", None),
        ("", None),
        ("各分頁", Font(name=FONT, bold=True, size=12)),
        ("病灶總表：每個病灶一列，包含偵測關鍵字、ICD-10-CM 代碼、對應的藥理類別與常用成分。", None),
        ("藥理類別明細：每個「病灶 × 藥理類別」一列，列出該類別在台灣有有效許可證的代表成分，請判斷該類別是否適合這個病灶。", None),
        ("ICD-10-CM 核對：目前代碼與健保署 2023 年版中文 ICD-10-CM 的比對結果與建議代碼。", None),
        ("病史風險提醒：病史出現某些關鍵字（或年齡 65 歲以上）時，系統會在相關藥品旁加註提醒。", None),
        ("缺少的病灶：請補充。", None),
        ("", None),
        ("名詞說明", Font(name=FONT, bold=True, size=12)),
        ("ATC：WHO 的藥品解剖治療化學分類。例：R05DA09（dextromethorphan）＝ R 呼吸系統 → R05 咳嗽感冒製劑（第 2 層）→ R05D 鎮咳劑（第 3 層）→ R05DA 鴉片生物鹼衍生物（第 4 層）→ 成分（第 5 層）。", None),
        ("系統目前以第 4 層比對：病灶對到的第 4 層類別底下的成分，才會出現在推薦中。「常用成分」會排在前面。", None),
        ("", None),
        ("資料來源", Font(name=FONT, bold=True, size=12)),
        ("藥品與成分：衛福部食藥署（TFDA）全部藥品許可證、ATC 碼資料集（data.gov.tw 9122、9119），只計入目前有效的許可證。", None),
        ("ICD-10-CM：衛福部中央健康保險署「2023年版_中文版ICD-10-CM/PCS（正式版）」（data.gov.tw 177507），健保自 2025-01-01 起使用。", None),
        ("ATC 類別名稱：WHO ATC 官方英文名稱；中文為本專案的參考翻譯。", None),
        ("病灶、關鍵字、類別對應、常用成分、風險提醒：開發者依一般用藥常識整理，待醫師審閱。", None),
    ]
    for i, (text, font) in enumerate(lines, 1):
        cell = ws.cell(i, 1, text)
        cell.font = font or Font(name=FONT)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # ---------- 病灶總表 ----------
    ws = wb.create_sheet("病灶總表")
    headers = [
        ("#", 5), ("病灶", 14), ("偵測關鍵字", 26), ("ICD-10-CM（目前）", 13), ("ICD 核對", 12),
        ("對應的藥理類別（ATC 第 4 層）", 46), ("所屬第 3 層", 34), ("常用成分（排在前面）", 34),
        ("避免推薦", 18), ("其他設定", 18), ("候選成分數", 10),
        ("審閱結果", 11), ("要新增的類別或成分", 24), ("要移除的類別或成分", 24), ("備註", 26),
    ]
    example = [
        "範例", "咳嗽", "咳嗽、乾咳", "R05", "需修正", "R05DA 鴉片生物鹼及衍生物", "R05D 鎮咳劑", "R05DA09 dextromethorphan",
        "", "", 44, "需修改", "R06AX 其他全身性抗組織胺應保留", "R05DA04 codeine（兒童不宜）", "（範例，非真實審閱）",
    ]
    rows = [example]
    for i, c in enumerate(conditions, 1):
        code = c["icd10"]
        status = icd_result(c["name"], icd.get(code, ("", ""))[0])
        others = []
        if c.get("chronic"):
            others.append("慢性病（候選依處方延續）")
        if c.get("exclude"):
            others.append("不算數的詞：" + "、".join(c["exclude"]))
        rows.append([
            i, c["name"], "、".join(c["aliases"]), code, status,
            "\n".join(atc_label(p) for p in c["atc"]),
            "\n".join(atc_label(p) for p in sorted({p[:4] for p in c["atc"]})),
            "\n".join(f"{p} {names.get(p, '')}" for p in c.get("first_line", [])),
            "\n".join(f"{p} {names.get(p, '')}" for p in c.get("avoid", [])),
            "\n".join(others),
            sum(1 for s in concepts if any(s[0].startswith(p) for p in c["atc"])),
            "", "", "", "",
        ])
    write_table(ws, headers, rows, input_cols={12, 13, 14, 15})
    for col in range(1, len(headers) + 1):
        ws.cell(2, col).fill = EXAMPLE_FILL
    for r in range(3, len(rows) + 2):
        status_cell = ws.cell(r, 5)
        status_cell.fill = OK_FILL if status_cell.value == "正確" else WARN_FILL
    add_choice(ws, 12, 2, len(rows) + 1, "同意,需修改,刪除")

    # ---------- 藥理類別明細 ----------
    ws = wb.create_sheet("藥理類別明細")
    headers = [
        ("病灶", 14), ("ATC 第 4 層", 9), ("類別名稱", 30), ("第 3 層", 30), ("有效成分數", 10),
        ("代表成分（依有效許可證數）", 50), ("此類別適合這個病灶嗎", 14), ("備註", 30),
    ]
    rows = [["咳嗽（範例）", "R05DA", "鴉片生物鹼及衍生物", "R05D 鎮咳劑", 5, "dextromethorphan、codeine…", "適合", "（範例）codeine 不建議用於 12 歲以下"]]
    for c in conditions:
        for p in c["atc"]:
            subs = substances(p)
            en, zh = ATC_NAMES.get(p, ("", ""))
            rows.append([
                c["name"], p, f"{zh}（{en}）" if en else "", atc_label(p[:4]), len(subs),
                "、".join(f"{substance_name(s[0], s[1])}（{s[2]}）" for s in subs[:8]) + ("…" if len(subs) > 8 else ""),
                "", "",
            ])
    write_table(ws, headers, rows, input_cols={7, 8})
    for col in range(1, len(headers) + 1):
        ws.cell(2, col).fill = EXAMPLE_FILL
    add_choice(ws, 7, 2, len(rows) + 1, "適合,不適合,視情況")

    # ---------- ICD-10-CM 核對 ----------
    ws = wb.create_sheet("ICD-10-CM 核對")
    headers = [
        ("病灶", 14), ("目前代碼", 11), ("2023 版中文名稱", 26), ("可申報", 8), ("核對結果", 12),
        ("建議代碼", 22), ("說明", 46), ("其他可選代碼", 46), ("醫師選用的代碼", 18),
    ]
    rows = []
    for c in conditions:
        code = c["icd10"]
        use, zh = icd.get(code, ("", "（ICD-10-CM 中沒有此代碼）"))
        suggest, note, alt = ICD_ADVICE.get(c["name"], (code, "正確", ""))
        result = icd_result(c["name"], use)
        rows.append([c["name"], code, zh, "是" if use == "1" else "否", result, suggest, note, alt, ""])
    write_table(ws, headers, rows, input_cols={9})
    for r in range(2, len(rows) + 2):
        result_cell = ws.cell(r, 5)
        result_cell.fill = OK_FILL if result_cell.value == "正確" else WARN_FILL

    # ---------- 病史風險提醒 ----------
    ws = wb.create_sheet("病史風險提醒")
    headers = [
        ("風險", 16), ("觸發條件（病史關鍵字或年齡）", 36), ("提醒的藥理類別", 46), ("系統顯示的提醒文字", 46),
        ("審閱結果", 11), ("要新增或移除的類別", 26), ("備註", 26),
    ]
    rows = []
    for f in flags:
        trigger = "、".join(f.get("keywords", [])) or f"年齡 ≥ {f.get('min_age')} 歲"
        rows.append([f["label"], trigger, "\n".join(atc_label(p) if p in ATC_NAMES else f"{p} {names.get(p, '')}" for p in f["atc"]), f["message"].replace("{matched}", "…").replace("{age}", "…"), "", "", ""])
    write_table(ws, headers, rows, input_cols={5, 6, 7})
    add_choice(ws, 5, 2, len(rows) + 1, "同意,需修改,刪除")

    # ---------- 缺少的病灶 ----------
    ws = wb.create_sheet("缺少的病灶")
    headers = [("病灶名稱", 18), ("常見說法（關鍵字）", 30), ("ICD-10-CM", 12), ("常用藥理類別或成分", 40), ("備註", 30)]
    rows = [["（範例）急性中耳炎", "耳朵痛、耳朵悶", "H66.90", "J01CR amoxicillin/clavulanate", "兒童常見"]] + [[""] * 5 for _ in range(15)]
    write_table(ws, headers, rows, input_cols={1, 2, 3, 4, 5})
    for col in range(1, 6):
        ws.cell(2, col).fill = EXAMPLE_FILL

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"完成：{OUT.relative_to(ROOT)}（{len(conditions)} 個病灶、{len(flags)} 項風險提醒）")


if __name__ == "__main__":
    main()
